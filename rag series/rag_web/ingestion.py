import hashlib
from pathlib import Path

import chromadb
import pandas as pd
import fitz

from docx import Document
from ollama import Client

from config import (
    DB_DIR,
    EMBED_MODEL,
    COLLECTION_NAME,
    CHUNK_SIZE,
    CHUNK_OVERLAP,
)

client = chromadb.PersistentClient(path=str(DB_DIR))
ollama = Client()


def extract_text(path):
    path = Path(path)
    ext = path.suffix.lower()
    units = []

    if ext == ".pdf":
        with fitz.open(path) as pdf:
            for page_no, page in enumerate(pdf, start=1):
                text = page.get_text().strip()
                if text:
                    units.append((text, page_no))

    elif ext == ".txt":
        units.append((
            path.read_text(encoding="utf-8"),
            None
        ))

    elif ext == ".docx":
        doc = Document(path)
        text = "\n".join(
            p.text for p in doc.paragraphs if p.text.strip()
        )

        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(
                    cell.text for cell in row.cells
                )

        units.append((text, None))

    elif ext in {".csv", ".xlsx"}:
        if ext == ".csv":
            df = pd.read_csv(path, dtype=str).fillna("")
        else:
            df = pd.read_excel(path, dtype=str).fillna("")

        for row_no, row in enumerate(
            df.to_dict(orient="records"), start=2
        ):
            text = "\n".join(
                f"{key}: {value}"
                for key, value in row.items()
                if str(value).strip()
            )

            if text.strip():
                units.append((text, row_no))

    return units


def chunk_text(text):
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == len(text):
            break

        start = end - CHUNK_OVERLAP

    return chunks


def ingest_files(file_paths, replace=True):
    if not file_paths:
        raise ValueError("Upload at least one document.")

    if replace:
        try:
            client.delete_collection(COLLECTION_NAME)
        except Exception:
            pass

        collection = client.create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"}
        )
    else:
        try:
            collection = client.get_collection(COLLECTION_NAME)
        except Exception:
            collection = client.create_collection(
                name=COLLECTION_NAME,
                metadata={"hnsw:space": "cosine"}
            )

    documents = []
    metadatas = []
    ids = []

    for raw_path in file_paths:
        path = Path(raw_path)
        units = extract_text(path)

        for text, location in units:
            for chunk_no, chunk in enumerate(chunk_text(text)):
                metadata = {
                    "source": path.name,
                    "file_type": path.suffix.lower().lstrip("."),
                }

                if path.suffix.lower() == ".pdf":
                    metadata["page"] = location
                elif path.suffix.lower() in {".csv", ".xlsx"}:
                    metadata["row"] = location

                chunk_id = hashlib.sha256(
                    f"{path.name}|{location}|{chunk_no}".encode()
                ).hexdigest()

                documents.append(chunk)
                metadatas.append(metadata)
                ids.append(chunk_id)

    if not documents:
        raise ValueError(
            "No readable text found. Scanned PDFs may require OCR."
        )

    # Embed in batches.
    for start in range(0, len(documents), 32):
        end = start + 32

        embeddings = ollama.embed(
            model=EMBED_MODEL,
            input=documents[start:end]
        )["embeddings"]

        collection.add(
            ids=ids[start:end],
            documents=documents[start:end],
            metadatas=metadatas[start:end],
            embeddings=embeddings
        )

    return collection.count()